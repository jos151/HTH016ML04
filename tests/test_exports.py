"""
Test Suite for Exports and Analytical Reports (Stage 17).
"""
import io
import openpyxl
import pandas as pd
import pytest
from backend.reports import generate_excel_report, generate_inventory_template


def test_excel_report_worksheets():
    """Verify that generate_excel_report produces a valid OpenPyXL workbook with all 8 sheets."""
    forecast_df = pd.DataFrame([
        {"store_id": "STORE_1", "sku_id": "SKU_01", "forecast_date": "2023-01-01", "predicted_units": 100.0, "baseline_units": 100.0},
        {"store_id": "STORE_2", "sku_id": "SKU_01", "forecast_date": "2023-01-01", "predicted_units": 50.0, "baseline_units": 50.0},
    ])
    allocation_df = pd.DataFrame([
        {"store_id": "STORE_1", "forecasted_demand": 100.0, "allocated_units": 80, "shortage": 20, "excess": 0},
        {"store_id": "STORE_2", "forecasted_demand": 50.0, "allocated_units": 40, "shortage": 10, "excess": 0},
    ])

    excel_bytes = generate_excel_report(
        forecast_df=forecast_df,
        allocation_df=allocation_df,
        summary_dict={"total_available_units": 120, "total_allocated_units": 120},
    )
    assert len(excel_bytes) > 0

    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    expected_sheets = [
        "Executive_Summary",
        "Forecast_Details",
        "Allocation_Details",
        "Store_Summary",
        "SKU_Summary",
        "Shortage_Risk",
        "Scenario_Assumptions",
        "Validation_Checks",
    ]
    for s in expected_sheets:
        assert s in wb.sheetnames, f"Worksheet {s} missing from Excel report."

    # Validate that Allocation_Details sheet contains data rows
    ws_alloc = wb["Allocation_Details"]
    assert ws_alloc.max_row >= 3  # Header + 2 stores


def test_inventory_template_csv():
    """Verify that generate_inventory_template creates a valid CSV template."""
    csv_bytes = generate_inventory_template(["SKU_01", "SKU_02"])
    assert len(csv_bytes) > 0
    df = pd.read_csv(io.BytesIO(csv_bytes))
    assert list(df.columns) == ["sku_id", "available_units", "warehouse_id", "safety_stock"]
    assert len(df) == 2
