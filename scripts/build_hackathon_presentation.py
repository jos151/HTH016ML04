"""
Builds the 6-slide executive PowerPoint presentation for the Final Hackathon Review.
Follows 16:9 widescreen layout, retail/supply-chain visual theme, strict typography,
integrated high-res visual assets, and embedded speaker notes.
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

OUTPUT_FILE = Path("D:/HTH016ML04/reports/Final_Hackathon_Review_Presentation.pptx")
ASSETS_DIR = Path("D:/HTH016ML04/reports/presentation_assets")
OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)

# Color Palette Constants
COLOR_NAVY_DARK = RGBColor(10, 25, 47)       # #0A192F
COLOR_NAVY_HEADER = RGBColor(15, 23, 42)     # #0F172A
COLOR_WHITE = RGBColor(255, 255, 255)        # #FFFFFF
COLOR_SLATE_BG = RGBColor(248, 250, 252)     # #F8FAFC
COLOR_SLATE_TEXT = RGBColor(71, 85, 105)     # #475569
COLOR_CYAN = RGBColor(2, 132, 199)           # #0284C7
COLOR_LIGHT_CYAN = RGBColor(56, 189, 248)    # #38BDF8
COLOR_ORANGE = RGBColor(234, 88, 12)         # #EA580C
COLOR_GREEN = RGBColor(5, 150, 105)          # #059669
COLOR_CARD_BG = RGBColor(241, 245, 249)      # #F1F5F9
COLOR_CARD_BORDER = RGBColor(203, 213, 225)  # #CBD5E1


def setup_presentation():
    prs = Presentation()
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs


def add_slide_header(slide, title_text, category_text=""):
    """Adds a standard enterprise header banner to a slide."""
    # Category Tracker
    if category_text:
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.35))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        tf_cat.margin_left = tf_cat.margin_top = tf_cat.margin_right = tf_cat.margin_bottom = 0
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = COLOR_CYAN

    # Main Slide Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.7), Inches(0.7))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.size = Pt(26)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_NAVY_HEADER


def set_speaker_notes(slide, notes_text):
    """Embeds concise speaker notes directly into the slide's notes view."""
    notes_slide = slide.notes_slide
    tf = notes_slide.notes_text_frame
    tf.text = notes_text


# ==============================================================================
# SLIDE 1: Title and Project Introduction
# ==============================================================================
def build_slide_1(prs):
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)

    # Dark Navy Background
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = COLOR_NAVY_DARK
    bg.line.fill.background()

    # Title & Subtitle Box
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.0), Inches(11.7), Inches(2.2))
    tf = title_box.text_frame
    tf.word_wrap = True

    p_badge = tf.paragraphs[0]
    p_badge.text = "RETAIL SUPPLY CHAIN INTELLIGENCE | HACKATHON 2026"
    p_badge.font.size = Pt(11)
    p_badge.font.bold = True
    p_badge.font.color.rgb = COLOR_LIGHT_CYAN
    p_badge.space_after = Pt(12)

    p_title = tf.add_paragraph()
    p_title.text = "Smart Demand Forecasting and Inventory Allocation"
    p_title.font.size = Pt(32)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_WHITE
    p_title.space_after = Pt(8)

    p_sub = tf.add_paragraph()
    p_sub.text = "Predict demand and distribute limited stock efficiently"
    p_sub.font.size = Pt(18)
    p_sub.font.color.rgb = RGBColor(148, 163, 184)
    p_sub.space_after = Pt(14)

    p_val = tf.add_paragraph()
    p_val.text = '"The system predicts store demand and recommends how limited inventory should be distributed."'
    p_val.font.size = Pt(13)
    p_val.font.italic = True
    p_val.font.color.rgb = RGBColor(226, 232, 240)

    # Visual Flow Image
    flow_img_path = ASSETS_DIR / "slide1_supply_chain_flow.png"
    if flow_img_path.exists():
        slide.shapes.add_picture(str(flow_img_path), Inches(1.5), Inches(3.4), width=Inches(10.33))

    # Footer Team Metadata
    footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.8), Inches(11.7), Inches(0.4))
    tf_f = footer_box.text_frame
    p_f = tf_f.paragraphs[0]
    p_f.text = "Team: Antigravity Solutions  |  Focus: Retail POS Analytics, Seasonality, & Constrained Rationing"
    p_f.font.size = Pt(10)
    p_f.font.color.rgb = RGBColor(100, 116, 139)

    # Speaker Notes (~30s)
    notes = (
        "Timing: ~30 seconds\n\n"
        "Good morning, judges and mentors. Today we are presenting our solution for Smart Demand Forecasting and "
        "Inventory Allocation.\n\n"
        "Retailers often have limited inventory in their central warehouse, but experience vastly different customer demand "
        "at each individual retail store. Our platform bridges this critical gap: it predicts future store-level demand and "
        "recommends how available stock should be distributed fairly and efficiently to eliminate stockouts and minimize excess inventory.\n\n"
        "Sources:\n"
        "- Visual on this slide was generated with Microsoft Copilot."
    )
    set_speaker_notes(slide, notes)


# ==============================================================================
# SLIDE 2: Problem and Business Need
# ==============================================================================
def build_slide_2(prs):
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)

    # Canvas Background
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = COLOR_SLATE_BG
    bg.line.fill.background()

    add_slide_header(slide, "The Problem We Are Solving", "Market Reality & Supply Scarcity")

    # Left Column: Key Problem Points (Max 5 bullets)
    left_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(5.8), Inches(4.0))
    tf_left = left_box.text_frame
    tf_left.word_wrap = True

    bullets = [
        "Each retail store experiences distinct customer demand patterns.",
        "Demand surges during weekends, calendar holidays, and promotions.",
        "Central warehouse stock is strictly limited and cannot meet peak demand.",
        "Poor allocation causes costly stockouts in high-demand stores.",
        "Oversupplying slower stores leads to trapped capital and surplus waste.",
    ]

    for i, b in enumerate(bullets):
        p = tf_left.paragraphs[0] if i == 0 else tf_left.add_paragraph()
        p.text = f"•  {b}"
        p.font.size = Pt(13)
        p.font.color.rgb = COLOR_NAVY_HEADER
        p.space_after = Pt(14)

    # Controlled Example Highlight Box
    callout = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.9), Inches(5.8), Inches(1.9))
    callout.fill.solid()
    callout.fill.fore_color.rgb = COLOR_CARD_BG
    callout.line.color.rgb = COLOR_CYAN
    callout.line.width = Pt(1.5)

    tf_c = callout.text_frame
    tf_c.word_wrap = True
    p_c1 = tf_c.paragraphs[0]
    p_c1.text = "THE CENTRAL CHALLENGE (Controlled Example)"
    p_c1.font.size = Pt(11)
    p_c1.font.bold = True
    p_c1.font.color.rgb = COLOR_CYAN
    p_c1.space_after = Pt(6)

    p_c2 = tf_c.add_paragraph()
    p_c2.text = "Total Store Demand = 1,200 units  |  Available Stock = 1,000 units\nSystem Shortage = 200 units (Shortfall)\nCentral Question: How should the available 1,000 units be distributed?"
    p_c2.font.size = Pt(11.5)
    p_c2.font.color.rgb = COLOR_NAVY_HEADER
    p_c2.font.bold = True

    # Right Column: Visual Comparison Gap Chart
    chart_path = ASSETS_DIR / "slide2_demand_gap_chart.png"
    if chart_path.exists():
        slide.shapes.add_picture(str(chart_path), Inches(6.9), Inches(1.6), width=Inches(5.7))

    # Speaker Notes (~60s)
    notes = (
        "Timing: ~60 seconds\n\n"
        "Let's look at the operational dilemma every inventory manager faces daily.\n\n"
        "In a retail network, no two stores are alike. Store A in a bustling city center might sell twice as fast as Store C in a suburb. "
        "On top of that, weekend rushes, promotional discounts, and holiday weeks introduce dramatic demand volatility.\n\n"
        "Meanwhile, upstream warehouse inventory is strictly finite. When aggregate demand across all stores exceeds available supply—like "
        "in our controlled scenario where three stores demand 1,200 units, but our warehouse only has 1,000 units—a shortage of 200 units is "
        "unavoidable.\n\n"
        "The wrong decision is to divide the stock equally: giving each store 333 units would leave Store A with a massive 167-unit stockout "
        "while oversupplying Store C beyond its needs. We need smart, demand-proportional allocation.\n\n"
        "Sources:\n"
        "- Visual on this slide was generated with Microsoft Copilot.\n"
        "- Data Source: Project-controlled demonstration scenario."
    )
    set_speaker_notes(slide, notes)


# ==============================================================================
# SLIDE 3: Our Solution and Workflow
# ==============================================================================
def build_slide_3(prs):
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)

    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = COLOR_SLATE_BG
    bg.line.fill.background()

    add_slide_header(slide, "How Our Solution Works", "End-to-End Decision Pipeline")

    # Workflow Graphic Banner across top/middle
    workflow_path = ASSETS_DIR / "slide3_solution_workflow.png"
    if workflow_path.exists():
        slide.shapes.add_picture(str(workflow_path), Inches(0.8), Inches(1.5), width=Inches(11.7))

    # 4 Supporting Points
    points_box = slide.shapes.add_textbox(Inches(0.8), Inches(4.7), Inches(11.7), Inches(1.6))
    tf_p = points_box.text_frame
    tf_p.word_wrap = True

    pts = [
        "Predict demand by store and product using historical rolling baselines and day-of-week seasonality.",
        "Adjust demand dynamically for targeted store promotions (+30%) and holiday weeks (+15%).",
        "Compare total predicted demand against available warehouse stock to identify systemic shortages.",
        "Allocate stock according to forecasted demand shares using discrete largest-remainder integer math.",
    ]

    for i, pt in enumerate(pts):
        p = tf_p.paragraphs[0] if i == 0 else tf_p.add_paragraph()
        p.text = f"•  {pt}"
        p.font.size = Pt(12.5)
        p.font.color.rgb = COLOR_NAVY_HEADER
        p.space_after = Pt(8)

    # Core Rule Banner at Bottom
    rule_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.2), Inches(11.7), Inches(0.8))
    rule_box.fill.solid()
    rule_box.fill.fore_color.rgb = RGBColor(239, 246, 255)
    rule_box.line.color.rgb = COLOR_CYAN
    rule_box.line.width = Pt(1.5)

    tf_r = rule_box.text_frame
    tf_r.word_wrap = True
    p_r = tf_r.paragraphs[0]
    p_r.text = 'Core Allocation Rule: "When supply is limited, each store receives stock based on its share of forecasted demand."'
    p_r.font.size = Pt(12)
    p_r.font.bold = True
    p_r.font.color.rgb = COLOR_CYAN
    p_r.alignment = PP_ALIGN.CENTER

    # Speaker Notes (~75s)
    notes = (
        "Timing: ~75 seconds\n\n"
        "Here is the six-stage architecture powering our solution, moving smoothly from raw data to actionable store decisions.\n\n"
        "First, we ingest historical point-of-sale transactions and clean the data across a complete Cartesian grid.\n\n"
        "Second, our forecasting engine models day-of-week seasonality on top of a 7-day rolling baseline, capturing cyclical shopping habits.\n\n"
        "Third, managers can simulate demand uplifts: promotions apply multiplier boosts to targeted stores, and holiday weeks scale demand "
        "across the network.\n\n"
        "Fourth, the manager inputs the actual warehouse inventory available for rationing.\n\n"
        "Fifth, our rationing engine applies Hare-Niemeyer largest-remainder integer math, or our optional PuLP mixed-integer linear programming "
        "solver, to eliminate fractional units and guarantee zero inventory wastage.\n\n"
        "Finally, store allocation schedules, stockout alerts, and multi-tab Excel reports are exported.\n\n"
        "The core principle is simple: every store receives whole units strictly according to its forecast demand share.\n\n"
        "Sources:\n"
        "- Visual on this slide was generated with Microsoft Copilot."
    )
    set_speaker_notes(slide, notes)


# ==============================================================================
# SLIDE 4: Web Application and Main Features
# ==============================================================================
def build_slide_4(prs):
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)

    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = COLOR_SLATE_BG
    bg.line.fill.background()

    add_slide_header(slide, "Web Application Features", "Interactive Decision-Support Platform")

    # Left Column: Features and User Workflow
    left_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(5.3), Inches(3.2))
    tf_l = left_box.text_frame
    tf_l.word_wrap = True

    p_fh = tf_l.paragraphs[0]
    p_fh.text = "CORE PLATFORM CAPABILITIES"
    p_fh.font.size = Pt(11)
    p_fh.font.bold = True
    p_fh.font.color.rgb = COLOR_CYAN
    p_fh.space_after = Pt(8)

    feats = [
        "Store and SKU demand forecasts with empirical uncertainty bounds.",
        "Inventory input and automated largest-remainder quota rationing.",
        "Real-time shortage, excess, and stockout risk alerts.",
        "What-if simulation for promotional (+30%) and holiday (+15%) uplifts.",
        "Side-by-side baseline vs scenario comparison metrics.",
        "One-click 8-worksheet Excel workbook & CSV report downloads.",
    ]
    for f in feats:
        p = tf_l.add_paragraph()
        p.text = f"•  {f}"
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_NAVY_HEADER
        p.space_after = Pt(6)

    # User Journey Steps Callout
    step_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.8), Inches(5.3), Inches(2.2))
    step_box.fill.solid()
    step_box.fill.fore_color.rgb = COLOR_CARD_BG
    step_box.line.color.rgb = COLOR_CARD_BORDER
    step_box.line.width = Pt(1)

    tf_s = step_box.text_frame
    tf_s.word_wrap = True
    p_sh = tf_s.paragraphs[0]
    p_sh.text = "SIMPLE 6-STEP DECISION FLOW"
    p_sh.font.size = Pt(10.5)
    p_sh.font.bold = True
    p_sh.font.color.rgb = COLOR_NAVY_HEADER
    p_sh.space_after = Pt(4)

    steps = [
        "1. Select stores and product SKUs",
        "2. Choose forecast horizon (7-30 days)",
        "3. Enter available warehouse inventory",
        "4. Run proportional or LP allocation",
        "5. Review store shortages & risk alerts",
        "6. Export full analytical Excel schedule",
    ]
    for s in steps:
        p = tf_s.add_paragraph()
        p.text = s
        p.font.size = Pt(9.5)
        p.font.color.rgb = COLOR_SLATE_TEXT

    # Right Column: Dashboard Mockup Image
    mockup_path = ASSETS_DIR / "slide4_dashboard_mockup.png"
    if mockup_path.exists():
        slide.shapes.add_picture(str(mockup_path), Inches(6.3), Inches(1.5), width=Inches(6.3))

    # Speaker Notes (~75s)
    notes = (
        "Timing: ~75 seconds\n\n"
        "To make this powerful forecasting and rationing logic usable for business operators, we packaged the complete pipeline "
        "into an intuitive, interactive Streamlit web dashboard backed by a high-performance FastAPI service.\n\n"
        "Here is the user workflow:\n"
        "First, planners select their active store and SKU parameters in the sidebar. They choose a forecast horizon—like 7 days—and "
        "enter the available warehouse inventory.\n\n"
        "With one click, the platform runs the forecasting and allocation engine. The Executive Dashboard instantly surfaces high-level "
        "KPI cards: total demand, total allocated units, system shortage, and overall network fill rate.\n\n"
        "Interactive bar charts allow planners to visually compare demand against allocated stock for every store, while prioritized alert "
        "cards flag high-risk stockout locations.\n\n"
        "Finally, operators can simulate promotional elasticity or download an 8-worksheet audit-ready Excel workbook with complete mathematical "
        "validation checks.\n\n"
        "Sources:\n"
        "- Source: Design Representation: Streamlit Web Dashboard."
    )
    set_speaker_notes(slide, notes)


# ==============================================================================
# SLIDE 5: Results and Controlled Demonstration
# ==============================================================================
def build_slide_5(prs):
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)

    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = COLOR_SLATE_BG
    bg.line.fill.background()

    add_slide_header(slide, "Allocation Result", "Controlled Demonstration & Mathematical Invariants")

    # Left Column: Table of Results & Validation Invariants
    left_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(5.6), Inches(5.5))
    tf_l = left_box.text_frame
    tf_l.word_wrap = True

    p_th = tf_l.paragraphs[0]
    p_th.text = "CONTROLLED DEMONSTRATION OUTCOME"
    p_th.font.size = Pt(11)
    p_th.font.bold = True
    p_th.font.color.rgb = COLOR_CYAN
    p_th.space_after = Pt(6)

    p_tbl = tf_l.add_paragraph()
    p_tbl.text = (
        "•  Store A:  Demand = 500  →  Allocated = 417  |  Shortage = 83  (83.4%)\n"
        "•  Store B:  Demand = 400  →  Allocated = 333  |  Shortage = 67  (83.3%)\n"
        "•  Store C:  Demand = 300  →  Allocated = 250  |  Shortage = 50  (83.3%)\n"
        "•  Summary: Demand = 1,200  |  Supply = 1,000  |  Allocated = 1,000  |  Shortage = 200"
    )
    p_tbl.font.size = Pt(10.5)
    p_tbl.font.color.rgb = COLOR_NAVY_HEADER
    p_tbl.space_after = Pt(12)

    p_vh = tf_l.add_paragraph()
    p_vh.text = "MATHEMATICAL INVARIANTS ENFORCED"
    p_vh.font.size = Pt(11)
    p_vh.font.bold = True
    p_vh.font.color.rgb = COLOR_GREEN
    p_vh.space_after = Pt(6)

    invs = [
        "Total allocation strictly respects warehouse supply: 417 + 333 + 250 = 1,000.",
        "No store receives more stock than its forecast demand (Ai <= Di).",
        "Allocation plus shortage exactly balances demand (Ai + Si = Di).",
        "Deterministic whole integer units (zero fractional cases).",
    ]
    for inv in invs:
        p = tf_l.add_paragraph()
        p.text = f"•  {inv}"
        p.font.size = Pt(10)
        p.font.color.rgb = COLOR_SLATE_TEXT
        p.space_after = Pt(4)

    # Actual Model Metrics Box
    metric_box = tf_l.add_paragraph()
    metric_box.text = (
        "\nFORECAST EVALUATION ON CHRONOLOGICAL HOLDOUT:\n"
        "MAE: 6.67 units  |  RMSE: 8.16 units  |  WAPE: 3.33%\n"
        "Forecast evaluation is performed using chronological test data."
    )
    metric_box.font.size = Pt(9.5)
    metric_box.font.bold = True
    metric_box.font.color.rgb = COLOR_NAVY_HEADER

    # Right Column: Grouped Bar Chart
    chart_path = ASSETS_DIR / "slide5_allocation_barchart.png"
    if chart_path.exists():
        slide.shapes.add_picture(str(chart_path), Inches(6.6), Inches(1.5), width=Inches(5.9))

    # Speaker Notes (~90s)
    notes = (
        "Timing: ~90 seconds\n\n"
        "Here are the concrete results from our controlled demonstration, validating how proportional largest-remainder "
        "math solves the scarcity dilemma.\n\n"
        "Looking at the table and grouped bar chart:\n"
        "Store A demanded 500 units and receives 417 units, leaving a shortage of 83.\n"
        "Store B demanded 400 units and receives 333 units, leaving a shortage of 67.\n"
        "Store C demanded 300 units and receives 250 units, leaving a shortage of 50.\n\n"
        "Notice what this achieves: every store achieves a balanced fill rate of approximately 83.3%. The system does not eliminate "
        "the 200-unit shortage—because available supply is fixed at 1,000—but it distributes that shortage fairly across stores in exact "
        "proportion to demand velocity.\n\n"
        "Furthermore, our automated test suite validates four mathematical conservation laws across every execution:\n"
        "1. Total allocation equals warehouse supply without inventory loss.\n"
        "2. No store receives more than its demand.\n"
        "3. Allocation plus shortage balances demand.\n"
        "4. All allocations are whole integers—essential for physical retail boxes.\n\n"
        "On historical out-of-sample data, our seasonal baseline achieves a Weighted Absolute Percentage Error of just 3.33%.\n\n"
        "Sources:\n"
        "- Source: Project-controlled test results."
    )
    set_speaker_notes(slide, notes)


# ==============================================================================
# SLIDE 6: Impact, Limitations, and Future Scope
# ==============================================================================
def build_slide_6(prs):
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)

    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = COLOR_SLATE_BG
    bg.line.fill.background()

    add_slide_header(slide, "Impact and Future Development", "Value Delivery & Product Roadmap")

    # Three-Column Framework Image in Center
    cards_path = ASSETS_DIR / "slide6_impact_summary.png"
    if cards_path.exists():
        slide.shapes.add_picture(str(cards_path), Inches(0.8), Inches(1.4), width=Inches(11.7))

    # Closing Value Banner
    banner = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.0), Inches(11.7), Inches(0.95))
    banner.fill.solid()
    banner.fill.fore_color.rgb = COLOR_NAVY_DARK
    banner.line.color.rgb = COLOR_CYAN
    banner.line.width = Pt(1.5)

    tf_b = banner.text_frame
    tf_b.word_wrap = True

    p_b1 = tf_b.paragraphs[0]
    p_b1.text = '"From prediction to allocation, our system helps retailers make practical inventory decisions."'
    p_b1.font.size = Pt(13)
    p_b1.font.bold = True
    p_b1.font.color.rgb = COLOR_WHITE
    p_b1.alignment = PP_ALIGN.CENTER
    p_b1.space_after = Pt(3)

    p_b2 = tf_b.add_paragraph()
    p_b2.text = "Thank You  |  Questions?"
    p_b2.font.size = Pt(11.5)
    p_b2.font.bold = True
    p_b2.font.color.rgb = COLOR_LIGHT_CYAN
    p_b2.alignment = PP_ALIGN.CENTER

    # Speaker Notes (~60s)
    notes = (
        "Timing: ~60 seconds\n\n"
        "To conclude, let's look at the operational impact, current boundaries, and future potential of our platform.\n\n"
        "Business Impact: Retailers gain immediate clarity. Instead of guessing replenishment orders, stores receive fair, "
        "demand-backed stock that minimizes lost sales while protecting against inventory holding costs.\n\n"
        "Current Limitations: Today, our baseline depends on historical point-of-sale transaction depth, models a single-warehouse "
        "tier, and requires available inventory to be user-provided.\n\n"
        "Future Roadmap: We plan to expand into multi-echelon network optimization, integrate advanced probabilistic ML models "
        "like DeepAR, automate supplier lead-time reordering, and deploy our native Android mobile application for on-the-go "
        "warehouse operations.\n\n"
        "From prediction to allocation, our system empowers retailers to make practical, data-driven inventory decisions.\n\n"
        "Thank you, and we welcome your questions.\n\n"
        "Sources:\n"
        "- Visual on this slide was generated with Microsoft Copilot."
    )
    set_speaker_notes(slide, notes)


def main():
    prs = setup_presentation()
    print("Building Slide 1: Title and Project Introduction...")
    build_slide_1(prs)
    print("Building Slide 2: Problem and Business Need...")
    build_slide_2(prs)
    print("Building Slide 3: Our Solution and Workflow...")
    build_slide_3(prs)
    print("Building Slide 4: Web Application Features...")
    build_slide_4(prs)
    print("Building Slide 5: Allocation Result...")
    build_slide_5(prs)
    print("Building Slide 6: Impact and Future Development...")
    build_slide_6(prs)

    prs.save(OUTPUT_FILE)
    print(f"PowerPoint successfully created and saved to: {OUTPUT_FILE}")
    print(f"Total slides created: {len(prs.slides)}")


if __name__ == "__main__":
    main()
